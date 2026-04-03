from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from .renderer import render_text_to_image


PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}\s]+)\s*\}\}")


@dataclass
class CellTarget:
    table_index: int
    row_index: int
    col_index: int


@dataclass
class RenderConfig:
    line_height: int = 56
    char_gap: int = 8


def _set_cell_text_with_image(cell, img_path: Path, width_mm: float = 40.0) -> None:
    # 清空原有段落后插入图片
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run()
    run.add_picture(str(img_path), width=Mm(width_mm))


def fill_word_table_by_placeholders(
    input_docx: str | Path,
    output_docx: str | Path,
    style_dir: str | Path,
    data: Dict[str, str],
    *,
    render: Optional[RenderConfig] = None,
    image_width_mm: float = 40.0,
    seed: Optional[int] = None,
) -> Path:
    """
    将 docx 表格内占位符 {{字段名}} 替换为手写字图片。
    """
    render = render or RenderConfig()

    input_docx = Path(input_docx)
    output_docx = Path(output_docx)
    document = Document(str(input_docx))

    with tempfile.TemporaryDirectory(prefix="handwrite_word_fill_") as tmp:
        tmpdir = Path(tmp)
        img_idx = 0

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    raw_text = "\n".join(p.text for p in cell.paragraphs).strip()
                    if not raw_text:
                        continue
                    m = PLACEHOLDER_RE.fullmatch(raw_text)
                    if not m:
                        continue
                    key = m.group(1)
                    if key not in data:
                        continue

                    value = str(data[key])
                    img_idx += 1
                    img_path = tmpdir / f"cell_{img_idx:04d}.png"
                    render_text_to_image(
                        style_dir=style_dir,
                        text=value,
                        out_path=img_path,
                        height=render.line_height,
                        char_gap=render.char_gap,
                        seed=seed,
                    )
                    _set_cell_text_with_image(cell, img_path, width_mm=image_width_mm)

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_docx))

    return output_docx


def fill_word_table_by_cell_map(
    input_docx: str | Path,
    output_docx: str | Path,
    style_dir: str | Path,
    cell_values: Dict[Tuple[int, int, int], str],
    *,
    render: Optional[RenderConfig] = None,
    image_width_mm: float = 40.0,
    seed: Optional[int] = None,
) -> Path:
    """
    按坐标写入表格，key=(table_idx,row_idx,col_idx), value=文本。
    """
    render = render or RenderConfig()

    input_docx = Path(input_docx)
    output_docx = Path(output_docx)
    document = Document(str(input_docx))

    with tempfile.TemporaryDirectory(prefix="handwrite_word_fill_") as tmp:
        tmpdir = Path(tmp)
        img_idx = 0

        for (t, r, c), value in cell_values.items():
            table = document.tables[t]
            cell = table.rows[r].cells[c]
            img_idx += 1
            img_path = tmpdir / f"coord_{img_idx:04d}.png"
            render_text_to_image(
                style_dir=style_dir,
                text=str(value),
                out_path=img_path,
                height=render.line_height,
                char_gap=render.char_gap,
                seed=seed,
            )
            _set_cell_text_with_image(cell, img_path, width_mm=image_width_mm)

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_docx))

    return output_docx
